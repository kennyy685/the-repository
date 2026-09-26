"""Apartment & commercial targets: bigger buildings inside recent hail zones, ranked by
hail at the building x job size x how recent x distance. One signed deal can equal many houses.

Owner names (public record) are added where the county publishes them openly
(Douglas / Omaha, Sarpy, Lancaster / Lincoln). Other counties: look up by the assessor link."""
import csv
import os
import time
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import numpy as np

from . import doors
from .geo import interp

DOUGLAS = "https://dcgis.org/server/rest/services/vector/Parcels_public/FeatureServer/0/query"
STATUS = ["Not contacted", "Called", "Left message", "Talked to manager", "Inspection set", "Damage found",
          "Claim filed", "Adjuster meeting", "Approved", "Job scheduled", "Done", "Lost"]
KIND_LABEL = {"multi": "Apartments / multi-family", "commercial": "Commercial", "industrial": "Industrial"}


def candidate_areas(conn, cfg, since, max_miles, min_hail, min_homes=800):
    """(storm day, town) pairs where hail >= min_hail touched a Nebraska town of min_homes+ homes."""
    towns = conn.execute("""SELECT name, state, lat, lon, aland_sqmi, hu, dist_mi FROM places
                            WHERE state='NE' AND COALESCE(hu,0) >= ? AND dist_mi <= ?""", (min_homes, max_miles)).fetchall()
    days = [r[0] for r in conn.execute("SELECT conv_day FROM swaths WHERE conv_day >= ? ORDER BY conv_day", (since,))]
    out = []
    for d in days:
        for t in towns:
            r = np.sqrt(max(t["aland_sqmi"] or 1, 0.5) / np.pi) + 1.0
            dlat, dlon = r / 69.05, r / (69.17 * np.cos(np.radians(t["lat"])))
            hit = conn.execute("""SELECT MAX(mesh_max_in) FROM nbhd_hits WHERE conv_day=? AND lat BETWEEN ? AND ?
                                  AND lon BETWEEN ? AND ?""", (d, t["lat"] - dlat, t["lat"] + dlat,
                                                              t["lon"] - dlon, t["lon"] + dlon)).fetchone()[0]
            if hit and hit >= min_hail:
                out.append((d, t["name"], r))
    return out


# County GIS services that publish owner names (public record). key = the county's parcel-id field.
OWNER_SOURCES = {
    "055": {"name": "Douglas", "url": "https://dcgis.org/server/rest/services/vector/Parcels_public/FeatureServer/0/query",
            "key": "PIN", "owner": ["OWNER_NAME"], "mail": ["ADDRESS1", "ADDRESS2", "OWNER_CITY", "OWNER_STAT", "OWNER_ZIP"],
            "desc": "BLDG_DESC", "link": "ASSESSOR"},
    "153": {"name": "Sarpy", "url": "https://geodata.sarpy.gov/arcgis/rest/services/Cadastral/LandRecordsSearch/MapServer/5/query",
            "key": "PARCELID", "owner": ["OWNERNME1", "OWNERNME2"], "mail": ["PSTLADDRESS", "PSTLCITY", "PSTLSTATE", "PSTLZIP5"],
            "desc": "CNVYNAME", "link": None},
    "109": {"name": "Lancaster", "url": "https://gis.lincoln.ne.gov/public/rest/services/Assessor/TaxParcels/FeatureServer/0/query",
            "key": "PARCELID", "owner": ["OWNERNME1", "OWNERNME2"], "mail": ["PSTLADDRESS", "PSTLCITY", "PSTLSTATE", "PSTLZIP5"],
            "desc": "CNVYNAME", "link": None},
}


def owners(session, county, parcel_ids):
    src = OWNER_SOURCES.get(county)
    out = {}
    if not src:
        return out
    ids = sorted({p for p in parcel_ids if p})
    fields = [src["key"]] + src["owner"] + src["mail"] + [src["desc"]] + ([src["link"]] if src["link"] else [])
    for k in range(0, len(ids), 150):
        where = f"{src['key']} IN (" + ",".join(f"'{p}'" for p in ids[k:k + 150]) + ")"
        try:
            r = session.post(src["url"], data={"where": where, "outFields": ",".join(fields), "returnGeometry": "false",
                                                "f": "json"}, timeout=90).json()
        except Exception:
            continue
        for f in r.get("features", []):
            a = f["attributes"]
            clean = lambda x: " ".join(str(x or "").split())
            out[a[src["key"]]] = {
                "owner": " & ".join(clean(a.get(o)) for o in src["owner"] if clean(a.get(o))),
                "owner_mail": ", ".join(clean(a.get(m)) for m in src["mail"] if clean(a.get(m))),
                "bldg_desc": clean(a.get(src["desc"])).title(), "bldgs": a.get("NUMBLDGS"),
                "assessor": clean(a.get(src["link"])) if src["link"] else ""}
    return out


def build(conn, cfg, session, since=None, max_miles=100, min_hail=1.0, kinds=("multi", "commercial"),
          budget_s=150, log=print):
    today = datetime.now(ZoneInfo(cfg["timezone"])).date()
    since = since or (today - timedelta(days=365)).isoformat()
    areas = candidate_areas(conn, cfg, since, max_miles, min_hail)
    log(f"  {len(areas)} storm-day x town combinations with {min_hail:g}\"+ hail since {since}")
    t0, best = time.monotonic(), {}
    done = 0
    for d, town, r in areas:
        if time.monotonic() - t0 > budget_s:
            log(f"  time budget reached after {done}/{len(areas)}; run again to continue (downloads are kept)")
            break
        bbox, _ = doors.area_bbox(conn, town, r)
        rows, _ = doors.score_buildings(conn, cfg, d, bbox, list(kinds), min_hail, session, log=lambda *a: None)
        done += 1
        for b in rows:
            b["storm_day"] = d
            cur = best.get(b["pid"])
            if cur is None or b["hail_in"] > cur["hail_in"] or (b["hail_in"] == cur["hail_in"] and d > cur["storm_day"]):
                b["n_storms"] = (cur or {}).get("n_storms", 0) + 1
                best[b["pid"]] = b
            else:
                cur["n_storms"] += 1
    sc = cfg["scoring"]
    out = []
    for b in best.values():
        days = (today - date.fromisoformat(b["storm_day"])).days
        size = min(1.0, max(0.2, ((b["imp_value"] or 0) / 2_000_000) ** 0.5))
        b["days_ago"] = days
        b["score"] = round(100 * interp(sc["size_curve"], b["hail_in"]) * interp(sc["recency_curve"], days) *
                           interp(sc["distance_curve"], b.get("dist_mi") or 0) * size, 1)
        out.append(b)
    out.sort(key=lambda b: -b["score"])
    own = {}
    for county in OWNER_SOURCES:
        ids = [b["parcel_id"] for b in out if b["county"] == county]
        if ids:
            own[county] = owners(session, county, ids)
            log(f"  owners from {OWNER_SOURCES[county]['name']} County: {len(own[county])}/{len(ids)}")
    for b in out:
        o = own.get(b["county"], {}).get(b["parcel_id"])
        b.update(o or {"owner": "", "owner_mail": "", "bldg_desc": "", "bldgs": None})
        if not b.get("assessor"):
            b["assessor"] = b.get("url") or ""
    return out, done, len(areas)


COLS = ["Rank", "Property address", "City", "Type", "Building", "Assessed building value ($)", "Year built",
        "Storm date", "Hail at building (in)", "Days ago", "Owner (public record)", "Owner mailing address",
        "Look up owner", "Score", "Status", "Contact / manager", "Phone", "Notes"]


def _row(k, b):
    return [k, b["address"], b["city"], KIND_LABEL.get(b["kind"], b["kind"]),
            (b.get("bldg_desc") or "") + (f" ({int(b['bldgs'])} bldgs)" if b.get("bldgs") and b["bldgs"] > 1 else ""),
            int(b["imp_value"]) if b["imp_value"] else None, b["year_built"] or None, b["storm_day"], b["hail_in"],
            b["days_ago"], b["owner"], b["owner_mail"], b["assessor"], b["score"], "Not contacted", "", "", ""]


def write_csv(out, cfg):
    """CSV only, standard library. Used by `refresh` on runners without openpyxl."""
    folder = os.path.join(cfg["paths"]["export"], "lists")
    os.makedirs(folder, exist_ok=True)
    stem = os.path.join(folder, f"apartments_commercial_{date.today().isoformat()}")
    with open(stem + ".csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        for k, b in enumerate(out, 1):
            w.writerow(_row(k, b))
    return stem + ".csv"


def write(out, cfg, title, subtitle):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    csv_path = write_csv(out, cfg)
    stem = csv_path[: -len(".csv")]
    F, head = "Arial", PatternFill("solid", fgColor="1C5CAB")
    yellow = PatternFill("solid", fgColor="FFF7D6")
    wb = Workbook()
    ws = wb.active
    ws.title = "Targets"
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font = Font(name=F, size=14, bold=True)
    ws["A2"].font = Font(name=F, size=9, color="52514E")
    for c, v in enumerate(COLS, 1):
        cell = ws.cell(row=4, column=c, value=v)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for k, b in enumerate(out, 1):
        for c, v in enumerate(_row(k, b), 1):
            cell = ws.cell(row=4 + k, column=c, value=v)
            cell.font = Font(name=F, size=10)
            name = COLS[c - 1]
            if name in ("Status", "Contact / manager", "Phone", "Notes"):
                cell.fill = yellow
            if name == "Assessed building value ($)":
                cell.number_format = '$#,##0'
            if name == "Hail at building (in)":
                cell.number_format = '0.00'
            if name == "Look up owner" and v:
                cell.hyperlink = v
                cell.value = "assessor page"
                cell.font = Font(name=F, size=10, color="1C5CAB", underline="single")
    last = 4 + max(len(out), 1)
    sc = get_column_letter(COLS.index("Status") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True)
    dv.add(f"{sc}5:{sc}{last}")
    ws.add_data_validation(dv)
    widths = [5, 26, 12, 20, 22, 14, 7, 11, 9, 7, 30, 34, 13, 7, 16, 20, 13, 30]
    for c, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = wdt
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(COLS))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    g = wb.create_sheet("How to use")
    g["A1"] = "Pipeline (updates as you change Status on the Targets sheet)"
    g["A1"].font = Font(name=F, size=12, bold=True)
    for k, s in enumerate(STATUS, 3):
        g.cell(row=k, column=1, value=s).font = Font(name=F)
        g.cell(row=k, column=2, value=f"=COUNTIF(Targets!{sc}5:{sc}{last},A{k})").font = Font(name=F)
    tips = [
        "HOW TO WORK THIS LIST",
        "1. Start at the top (highest score = most hail on the biggest, most recent buildings).",
        "2. Owner (public record) is filled in for Douglas (Omaha), Sarpy and Lancaster (Lincoln) counties. Elsewhere, look up the address on the county assessor website.",
        "3. Owners are often LLCs. Find the manager: Google the property name, call the leasing office and ask who handles "
        "roofing/exterior, or look up the LLC's registered agent at the Nebraska Secretary of State business search.",
        "4. Pitch: 'Your property was in the hail path on <storm date> (about <hail> inch hail per NOAA radar). We're a local "
        f"{cfg['home']['name'].split(',')[0].strip()} siding & roofing company; we do apartments already. Can we do a free inspection and a written report?'",
        "5. Fill in the yellow columns; the pipeline counts above update themselves.",
        "Example: Status = Inspection set | Contact = Jen (property manager) | Phone = 402-555-0100 | Notes = walk roof Thu 9am",
        "Commercial claims and property manager approvals take longer than homes. Keep following up.",
        "Hail at building = NOAA MRMS radar corrected with ground reports (about 1 km detail). Always inspect before quoting.",
        "Sources: Nebraska Statewide Parcels; Douglas County GIS (owner public record); NOAA MRMS; NWS storm reports.",
    ]
    for k, t in enumerate(tips):
        c = g.cell(row=3 + len(STATUS) + 2 + k, column=1, value=t)
        c.font = Font(name=F, bold=(k == 0), size=10)
    g.column_dimensions["A"].width = 120
    g.column_dimensions["B"].width = 8
    wb.move_sheet("How to use", offset=-1)
    wb.save(stem + ".xlsx")
    return stem + ".xlsx", stem + ".csv"

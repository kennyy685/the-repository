"""Offline regression tests for the 2026-09-26 review fixes: wind units + retries, refresh ordering,
hud neighborhood hail, door-sheet labels, and the fresh-database places crash (T19)."""
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta, timezone
from unittest import mock

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import hh  # noqa: E402
from hailhunter import commercial, db, doors, hud, ingest, mrms, nbhd, wind  # noqa: E402
from hailhunter import config as C  # noqa: E402
from hailhunter.http import Fetcher  # noqa: E402
from hailhunter.models import iso  # noqa: E402
from hailhunter.sources import places  # noqa: E402

BG = "310539640003"
RING = [[-96.60, 41.50], [-96.50, 41.50], [-96.50, 41.40], [-96.60, 41.40], [-96.60, 41.50]]


def _lsr(features):
    return json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-96.5, 41.44]},
         "properties": {"valid": "2026-08-17T22:10:00Z", "source": "Trained Spotter", "city": "Fremont",
                        "county": "Dodge", "st": "NE", "wfo": "OAX", "remark": "", **p}}
        for p in features]}).encode()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.cfg = C.load(os.path.join(self.tmp, "none.json"))
        self.cfg["paths"] = {k: os.path.join(self.tmp, k) for k in ("db", "cache", "export")}
        os.makedirs(self.cfg["paths"]["export"])
        self.conn = db.connect(self.cfg["paths"]["db"])
        self.f = Fetcher(self.cfg["paths"]["cache"], offline=True)

    def tearDown(self):
        self.conn.close()
        shutil.rmtree(self.tmp)


class Wind(Base):
    def test_land_gusts_are_mph_marine_are_knots(self):
        obs = wind.parse(_lsr([
            {"typetext": "TSTM WND GST", "magnitude": 60, "unit": "MPH"},
            {"typetext": "MARINE TSTM WND", "magnitude": 50, "unit": "KT"},
            {"typetext": "TSTM WND DMG", "magnitude": None, "unit": None, "remark": "Roof blown off"},
        ]), self.cfg)
        land, marine, dmg = obs
        self.assertEqual((land["speed_mph"], land["speed_kt"]), (60, 52.1))     # was 69.0 mph before the fix
        self.assertEqual((marine["speed_mph"], marine["speed_kt"]), (57.5, 50))
        self.assertEqual((dmg["speed_mph"], dmg["speed_kt"], dmg["report_kind"]), (None, None, "damage"))

    def test_failed_window_is_retried_next_run(self):
        s = datetime(2026, 6, 1, 12, tzinfo=timezone.utc)
        e = datetime(2026, 6, 10, 12, tzinfo=timezone.utc)
        _, st = wind.ingest(self.conn, self.f, self.cfg, s, e, log=lambda *a: None)   # nothing cached: all fail
        self.assertTrue(st["errors"])
        run = self.conn.execute("SELECT status, reached_utc FROM ingest_runs WHERE source='wind'").fetchone()
        self.assertEqual(run["status"], "failed")
        self.assertEqual(run["reached_utc"], iso(s))          # resume from the failure, not from the end

    def test_old_rows_repaired_on_connect(self):
        rows = [("w1", "TSTM WND GST", 60.0, 69.0), ("w2", "MARINE TSTM WND", 50.0, 57.5)]
        for uid, tt, kt, mph in rows:
            self.conn.execute("""INSERT INTO wind_obs (uid, source, valid_utc, conv_day, lat, lon, speed_kt,
                                 speed_mph, extra) VALUES (?, 'lsr', '2026-08-17T22:10:00Z', '2026-08-17', 41.4,
                                 -96.5, ?, ?, ?)""", (uid, kt, mph, json.dumps({"typetext": tt})))
        self.conn.execute("DELETE FROM meta WHERE key='wind_mph_fix'")      # a database from before the fix
        self.conn.commit()
        self.conn.close()
        self.conn = db.connect(self.cfg["paths"]["db"])
        got = {r[0]: (r[1], r[2]) for r in self.conn.execute("SELECT uid, speed_kt, speed_mph FROM wind_obs")}
        self.assertEqual(got["w1"], (52.1, 60.0))
        self.assertEqual(got["w2"], (50.0, 57.5))                           # marine untouched
        self.conn.close()
        self.conn = db.connect(self.cfg["paths"]["db"])                     # runs once only
        self.assertEqual(self.conn.execute("SELECT speed_mph FROM wind_obs WHERE uid='w1'").fetchone()[0], 60.0)


class Refresh(Base):
    """refresh() with the network pieces stubbed out, to check what gets measured and in which order."""

    def setUp(self):
        super().setUp()
        self.conn.execute("INSERT INTO bgs VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                          (BG, "NE", "053", "964000", "3", 41.45, -96.55, 2.0, 3.0, -96.60, 41.40, -96.50, 41.50,
                           json.dumps([RING]), "T01", "Testville, NE center [9640-3]"))
        self.conn.execute(f"INSERT INTO acs VALUES ('{BG}','bg',400,380,338,42,1950,180000,'2024')")
        self.conn.commit()

    def _swath(self, day, inches=1.0):
        arr = np.full((60, 60), round(inches * 25.4 * 10), np.uint16)
        meta = {"lat0": 41.6, "lon0": -96.7, "dlat": 0.01, "dlon": 0.01, "shape": [60, 60], "complete": True}
        with open(mrms.grid_path(self.cfg, day), "wb") as f:
            np.savez_compressed(f, mesh=arr, meta=json.dumps(meta))
        self.conn.execute("INSERT INTO swaths VALUES (?,?,?,1,?,0,0,0,0,?)", (day, "k", "t", inches, "t"))
        self.conn.commit()

    def _cal(self, factor, age_days=0):
        return {"factor": factor, "pairs": 40, "p25": factor, "p75": factor, "days": 1,
                "computed_utc": iso(datetime.now(timezone.utc) - timedelta(days=age_days))}

    def _refresh(self, touched=(), done=(), new_factor=None):
        def fake_calibrate(conn, cfg, log=print):
            res = self._cal(new_factor)
            db.set_meta(conn, "mesh_calibration", res)
            return res
        with mock.patch.object(ingest, "run", return_value=(set(touched), {})), \
                mock.patch.object(wind, "ingest") as w, \
                mock.patch.object(mrms, "run", return_value=(list(done), 0, [])), \
                mock.patch.object(nbhd, "calibrate", side_effect=fake_calibrate) as cal, \
                mock.patch.object(hh, "pick_door_lists", return_value=[]), \
                mock.patch.object(commercial, "build", return_value=([], 0, 0)), \
                mock.patch.object(commercial, "write_csv"):
            hh.refresh(self.conn, self.f, self.cfg, log=lambda *a: None)
        return w, cal

    def _hail(self, day):
        r = self.conn.execute("SELECT hail_in FROM nbhd_hits WHERE conv_day=? AND geoid=?", (day, BG)).fetchone()
        return r and r[0]

    def test_first_run_calibrates_before_measuring_and_pulls_wind(self):
        self._swath("2026-08-17")
        w, cal = self._refresh(done=["2026-08-17"], new_factor=1.5)
        w.assert_called_once()                                              # REQ-4
        cal.assert_called_once()
        self.assertAlmostEqual(self._hail("2026-08-17"), 1.5, places=2)    # was measured at x1.0 before

    def test_late_ground_reports_remeasure_old_day(self):
        self._swath("2026-07-01")
        db.set_meta(self.conn, "mesh_calibration", self._cal(1.0))
        self._refresh(touched=["2026-07-01"], done=[])                     # map already stored, reports changed
        self.assertIsNotNone(self._hail("2026-07-01"))

    def test_monthly_recalibration_remeasures_every_day(self):
        self._swath("2026-07-01")
        self._swath("2026-08-17")
        db.set_meta(self.conn, "mesh_calibration", self._cal(1.0, age_days=40))
        self._refresh(done=["2026-08-17"], new_factor=1.2)
        self.assertAlmostEqual(self._hail("2026-07-01"), 1.2, places=2)    # old day picked up the new factor

    def test_fresh_calibration_is_not_redone(self):
        self._swath("2026-08-17")
        db.set_meta(self.conn, "mesh_calibration", self._cal(1.0, age_days=3))
        _, cal = self._refresh(done=["2026-08-17"], new_factor=1.5)
        cal.assert_not_called()


class Outputs(Base):
    def test_hud_neighborhood_hail_is_max_and_avg_is_typical(self):
        day = (date.today() - timedelta(days=10)).isoformat()
        self.conn.execute("""INSERT INTO nbhd_hits (conv_day, geoid, label, place_name, state, lat, lon, dist_mi, hu,
                             owner_share, med_year, mesh_max_in, mesh_p75_in, hail_in, frac_ge_1, score)
                             VALUES (?, ?, 'Testville', 'Testville', 'NE', 41.45, -96.55, 3, 400, 0.8, 1978,
                                     1.8, 1.3, 1.3, 0.9, 50)""", (day, BG))
        n = hud._neighborhoods(self.conn, self.cfg, (date.today() - timedelta(days=365)).isoformat())[0]
        self.assertEqual((n["hail"], n["hail_avg"]), (1.8, 1.3))

    @unittest.skipUnless(__import__("importlib").util.find_spec("openpyxl"), "openpyxl not installed")
    def test_door_sheet_uses_current_labels_and_per_town_permit(self):
        from openpyxl import load_workbook
        h = {"address": "1 Main St", "city": "Columbus", "zip": "68601", "kind": "single", "year_built": 1990,
             "sqft": 1200, "hail_in": 1.5, "flags": "", "score": 50, "lat": 41.4, "lon": -97.3}
        turfs = [{"stops": [h], "doors": 1, "avg_hail": 1.5, "value": 50, "streets": "Main St"}]
        path = os.path.join(self.tmp, "t.xlsx")
        doors.write_xlsx(path, turfs, "t", "s")
        ws = load_workbook(path)["Summary"]
        self.assertIn('"Interested"', ws["G5"].value)
        self.assertNotIn('"Lead"', ws["G5"].value)
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        self.assertNotIn("Fremont solicitor", text)
        self.assertIn("each town needs its own", text)


class FreshDatabase(Base):
    def test_places_load_on_empty_database(self):
        """T19: an 8-value INSERT crashed on a brand-new database once `places` grew a 9th column."""
        gaz = "USPS\tGEOID\tNAME\tALAND_SQMI\tINTPTLAT\tINTPTLONG\nNE\t3117670\tFremont city\t9.1\t41.44\t-96.49\n"
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("gaz.txt", gaz)
        self.f.put(places.GAZ[0], buf.getvalue())
        n = places.load(self.conn, self.f, self.cfg, log=lambda *a: None)   # Census population API: offline, skipped
        self.assertEqual(n, 1)
        self.assertEqual(self.conn.execute("SELECT name FROM places").fetchone()[0], "Fremont")


if __name__ == "__main__":
    unittest.main()
